from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs
import json
import os
from pathlib import Path
import base64
import io
from datetime import date, datetime, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / 'data'
DATA_DIR.mkdir(exist_ok=True)
STORAGE_FILE = DATA_DIR / 'storage.json'

if not STORAGE_FILE.exists():
    STORAGE_FILE.write_text('{}', encoding='utf-8')

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/storage':
            self.handle_storage_get(parsed)
            return
        super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/storage':
            self.handle_storage_post(parsed)
            return
        if parsed.path == '/api/export-excel':
            self.handle_export_excel(parsed)
            return
        super().do_POST()

    def do_DELETE(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/storage':
            self.handle_storage_delete(parsed)
            return
        super().do_DELETE()

    def handle_storage_get(self, parsed):
        params = parse_qs(parsed.query)
        key = params.get('key', [''])[0]
        data = self.load_storage()
        value = data.get(key)
        payload = {'value': value} if value is not None else {}
        self.send_json(payload)

    def handle_storage_post(self, parsed):
        length = int(self.headers.get('Content-Length', '0'))
        body = self.rfile.read(length).decode('utf-8') if length else '{}'
        payload = {}
        try:
            payload = json.loads(body)
        except json.JSONDecodeError:
            try:
                payload = {k: v[0] if len(v) == 1 else v for k, v in parse_qs(body).items()}
            except Exception:
                payload = {}

        key = payload.get('key')
        value = payload.get('value')
        if not key:
            self.send_json({'ok': False}, status=400)
            return

        data = self.load_storage()
        data[key] = value
        self.save_storage(data)
        self.send_json({'ok': True})

    def handle_export_excel(self, parsed):
        if Workbook is None:
            self.send_json({'error': 'OpenPyXL no está disponible en el servidor.'}, status=500)
            return

        length = int(self.headers.get('Content-Length', '0'))
        try:
            payload = json.loads(self.rfile.read(length).decode('utf-8'))
            issues = payload.get('issues', [])
            images = payload.get('images', {})
            output = self.build_excel(issues, images, payload.get('scope', ''))
        except Exception as exc:
            self.send_json({'error': f'No se pudo generar el Excel: {exc}'}, status=500)
            return

        body = output.getvalue()
        filename = f"issues-{__import__('datetime').date.today().isoformat()}.xlsx"
        self.send_response(200)
        self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.send_header('Content-Disposition', f'attachment; filename="{filename}"')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def build_excel(self, issues, images, scope=''):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Bitácora'
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False

        headers = ['Test ID', 'Título del hallazgo', 'Módulo', 'Descripción del hallazgo',
                   'Estado', 'Severidad', 'Prioridad', 'Asignado a', 'Fecha registro',
                   'Antigüedad (días)', 'Fecha cierre', 'Última actualización',
                   'Comentarios / trazabilidad', 'Imágenes', 'Evidencia 1',
                   'Evidencia 2', 'Evidencia 3', 'Evidencia 4', 'Evidencia 5']
        sheet.append(headers)

        header_fill = PatternFill('solid', fgColor='FF1F3A5F')
        thin_border = Border(*(Side(style='thin', color='D4DEE7') for _ in range(4)))
        body_font = Font(name='Arial', size=10)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        sheet.row_dimensions[1].height = 26

        width_map = [14, 30, 22, 44, 18, 14, 14, 22, 16, 16, 16, 18, 46, 12, 24, 24, 24, 24, 24]
        for index, width in enumerate(width_map, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        status_labels = {'Open': 'Open', 'Blocked': 'Blocked', 'ReadyRetest': 'Ready for Retest',
                         'InProgress': 'In Progress', 'Closed': 'Closed', 'ReOpen': 'Re-Open'}
        for issue_index, issue in enumerate(issues, 2):
            created = self.parse_report_date(issue.get('fecha') or issue.get('createdAt'))
            solution_date = self.report_solution_date(issue)
            update_dates = [
                self.parse_report_date(item.get('createdAt') or item.get('date'))
                for item in issue.get('novedades', []) if isinstance(item, dict)
            ]
            last_update = max((value for value in update_dates if value), default=None)
            status = issue.get('status', '')
            end_date = solution_date if status == 'Closed' else date.today()
            age = max(0, (end_date - created).days) if created and end_date else None
            trace = '\n\n'.join(
                item if isinstance(item, str) else
                f"[{self.report_date_text(item.get('createdAt') or item.get('date')) or 'Sin fecha'}] "
                f"{item.get('type', 'Novedad')}: {item.get('text', '')}"
                for item in issue.get('novedades', []) if isinstance(item, (str, dict))
            )
            issue_images = images.get(str(issue.get('id')), [])[:5]
            values = [issue.get('testid', ''), issue.get('formulario', ''), issue.get('modulo', ''),
                      issue.get('obs', ''), status_labels.get(status, status), issue.get('severity', ''),
                      issue.get('priority', ''), issue.get('assigned', ''),
                      created or '', age if age is not None else '', solution_date or '',
                      last_update or '', trace, len(issue_images), '', '', '', '', '']
            sheet.append(values)
            sheet.row_dimensions[issue_index].height = 98
            for cell in sheet[issue_index]:
                cell.font = body_font
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.border = thin_border
            for column in (9, 11, 12):
                sheet.cell(issue_index, column).number_format = 'dd/mm/yyyy'

            for image_index, data_url in enumerate(issue_images):
                image = self.decode_excel_image(data_url)
                if image:
                    scale = min(110 / image.width, 82 / image.height)
                    image.width = int(image.width * scale)
                    image.height = int(image.height * scale)
                    sheet.add_image(image, f'{get_column_letter(15 + image_index)}{issue_index}')

        sheet.auto_filter.ref = f'A1:S{len(issues) + 1}'

        metrics = workbook.create_sheet('Resumen')
        metrics.sheet_view.showGridLines = False
        metrics.freeze_panes = 'A5'
        self.populate_metrics(metrics, issues, images, scope, status_labels)

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @staticmethod
    def parse_report_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if not value:
            return None
        text = str(value).strip()
        for date_format in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%Y-%m-%dT%H:%M:%S.%fZ',
                            '%Y-%m-%dT%H:%M:%S.%f', '%Y-%m-%dT%H:%M:%S'):
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(text.replace('Z', '+00:00')).date()
        except ValueError:
            return None

    @staticmethod
    def report_date_text(value):
        parsed = Handler.parse_report_date(value)
        return parsed.strftime('%d/%m/%Y') if parsed else ''

    def report_solution_date(self, issue):
        dates = [
            self.parse_report_date(item.get('createdAt') or item.get('date'))
            for item in issue.get('novedades', [])
            if isinstance(item, dict) and item.get('type') == 'Solución'
        ]
        return max((parsed for parsed in dates if parsed), default=None)

    def decode_excel_image(self, data_url):
        prefix, separator, payload = data_url.partition(',')
        if not separator or not prefix.startswith('data:image/'):
            return None
        raw = base64.b64decode(payload)
        image = Image(io.BytesIO(raw))
        image.width = 110
        image.height = 82
        return image

    def populate_metrics(self, metrics, issues, images, scope, status_labels):
        total = len(issues)
        closed_issues = [issue for issue in issues if issue.get('status') == 'Closed']
        open_issues = [issue for issue in issues if issue.get('status') != 'Closed']
        status_counts = {status: sum(issue.get('status') == status for issue in issues)
                         for status in status_labels}
        severity_order = ['Critical', 'Crítica', 'Alta', 'High', 'Media', 'Medium', 'Baja', 'Low']
        severities = set(['Alta', 'Media', 'Baja'])
        severities.update(issue.get('severity') for issue in issues if issue.get('severity'))
        severities = sorted(severities, key=lambda value: (
            severity_order.index(value) if value in severity_order else len(severity_order), value.lower()
        ))
        modules = sorted({str(issue.get('modulo') or 'Sin definir') for issue in issues}, key=str.lower)
        owners = sorted({str(issue.get('assigned') or 'Sin asignar') for issue in issues}, key=str.lower)
        evidence_count = sum(len(images.get(str(issue.get('id')), [])) for issue in issues)

        metrics.append(['Resumen de issues / Bitácora QA'])
        metrics.append([f'Alcance: {total} issues (según los filtros aplicados) · {scope or "Todos los registros"} · {date.today().strftime("%d/%m/%Y")}'])
        metrics.append([])

        def add_section(title, headers, rows, percentage_columns=()):
            metrics.append([title])
            metrics.append(headers)
            for row in rows:
                metrics.append(row)
                if row and row[0] == 'Total':
                    for cell in metrics[metrics.max_row]:
                        cell.fill = PatternFill('solid', fgColor='EAF0F6')
                        cell.font = Font(name='Arial', size=10, bold=True)
                for column in percentage_columns:
                    metrics.cell(metrics.max_row, column).number_format = '0.0%'
            metrics.append([])

        status_rows = [[status_labels[status], count, count / total if total else 0]
                       for status, count in status_counts.items()]
        status_rows.append(['Total', total, 1 if total else 0])
        add_section('Issues por estado', ['Estado', 'Cantidad', '% del total'], status_rows, (3,))

        severity_rows = []
        for severity in severities:
            matches = [issue for issue in issues if issue.get('severity') == severity]
            severity_rows.append([severity, sum(issue.get('status') != 'Closed' for issue in matches),
                                  sum(issue.get('status') == 'Closed' for issue in matches), len(matches)])
        add_section('Issues por severidad', ['Severidad', 'Abiertos', 'Cerrados', 'Total'], severity_rows)

        module_rows = []
        for module in modules:
            matches = [issue for issue in issues if (issue.get('modulo') or 'Sin definir') == module]
            module_rows.append([module, sum(issue.get('status') != 'Closed' for issue in matches),
                                sum(issue.get('status') == 'Closed' for issue in matches), len(matches)])
        add_section('Issues por módulo', ['Módulo', 'Abiertos', 'Cerrados', 'Total'], module_rows)

        week_rows = []
        current_monday = date.today() - timedelta(days=date.today().weekday())
        for offset in range(7, -1, -1):
            start = current_monday - timedelta(weeks=offset)
            end = start + timedelta(days=7)
            registered = sum(
                (created := self.parse_report_date(issue.get('fecha') or issue.get('createdAt'))) is not None
                and start <= created < end for issue in issues
            )
            resolved = sum(
                issue.get('status') == 'Closed' and
                (closed := self.report_solution_date(issue)) is not None and start <= closed < end
                for issue in issues
            )
            week_rows.append([start.strftime('%d %b'), registered, resolved])
        add_section('Tendencia semanal (últimas 8 semanas)',
                    ['Semana del', 'Registrados', 'Cerrados'], week_rows)

        owner_rows = []
        for owner in owners:
            matches = [issue for issue in issues if (issue.get('assigned') or 'Sin asignar') == owner]
            owner_rows.append([owner, sum(issue.get('status') != 'Closed' for issue in matches),
                               sum(issue.get('status') == 'Closed' for issue in matches), len(matches)])
        add_section('Issues por responsable', ['Responsable', 'Abiertos', 'Cerrados', 'Total'], owner_rows)

        open_ages = []
        for issue in open_issues:
            created = self.parse_report_date(issue.get('fecha') or issue.get('createdAt'))
            if created:
                open_ages.append(max(0, (date.today() - created).days))
        closure_durations = []
        for issue in closed_issues:
            created = self.parse_report_date(issue.get('fecha') or issue.get('createdAt'))
            resolved = self.report_solution_date(issue)
            if created and resolved:
                closure_durations.append(max(0, (resolved - created).days))
        reopened = sum(issue.get('status') == 'ReOpen' for issue in issues)
        quality_rows = [
            ['Total de issues', total],
            ['Issues cerrados', len(closed_issues)],
            ['Issues abiertos', len(open_issues)],
            ['Resolución (% cerrados sobre el total)', len(closed_issues) / total if total else 0],
            ['Evidencias adjuntas', evidence_count],
            ['Tiempo promedio de cierre (días)', sum(closure_durations) / len(closure_durations) if closure_durations else 'Sin datos'],
            ['Antigüedad promedio de issues abiertos (días)', sum(open_ages) / len(open_ages) if open_ages else 'Sin datos'],
            ['Issue abierto más antiguo (días)', max(open_ages) if open_ages else 'Sin datos'],
            ['Issues abiertos sin responsable', sum(not issue.get('assigned') for issue in open_issues)],
            ['Issues en estado Re-Open', reopened],
            ['Tasa de reapertura (%)', reopened / total if total else 0]
        ]
        add_section('Indicadores de calidad y gestión', ['Indicador', 'Valor'], quality_rows, (2,))

        metrics.column_dimensions['A'].width = 46
        metrics.column_dimensions['B'].width = 20
        metrics.column_dimensions['C'].width = 18
        metrics.column_dimensions['D'].width = 18
        border = Border(*(Side(style='thin', color='D4DEE7') for _ in range(4)))
        navy = PatternFill('solid', fgColor='FF1F3A5F')
        pale = PatternFill('solid', fgColor='FFEAF0F6')
        for row_index, row in enumerate(metrics.iter_rows(), 1):
            title = row_index == 1 or (row[0].value and str(row[0].value).startswith(('Issues por ', 'Tendencia semanal', 'Indicadores de ')))
            header = row_index > 1 and row[0].value in ('Estado', 'Severidad', 'Módulo', 'Semana del', 'Responsable', 'Indicador')
            for cell in row:
                cell.font = Font(name='Arial', size=10)
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if title:
                    cell.fill = navy
                    cell.font = Font(name='Arial', size=14 if row_index == 1 else 10, bold=True, color='FFFFFF')
                elif header:
                    cell.fill = navy
                    cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                elif cell.value is not None and row_index % 2 == 0:
                    cell.fill = pale
            if row_index == 1:
                metrics.row_dimensions[row_index].height = 30
            elif title:
                metrics.row_dimensions[row_index].height = 24
            elif header:
                metrics.row_dimensions[row_index].height = 22

    def handle_storage_delete(self, parsed):
        params = parse_qs(parsed.query)
        key = params.get('key', [''])[0]
        if not key:
            self.send_json({'ok': False}, status=400)
            return

        data = self.load_storage()
        data.pop(key, None)
        self.save_storage(data)
        self.send_json({'ok': True})

    def load_storage(self):
        try:
            with STORAGE_FILE.open('r', encoding='utf-8') as fh:
                return json.load(fh)
        except Exception:
            return {}

    def save_storage(self, data):
        with STORAGE_FILE.open('w', encoding='utf-8') as fh:
            json.dump(data, fh, ensure_ascii=False, indent=2)

    def send_json(self, payload, status=200):
        body = json.dumps(payload).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', '8000'))
    server = ThreadingHTTPServer(('0.0.0.0', port), Handler)
    print(f'Serving at http://127.0.0.1:{port}')
    server.serve_forever()
